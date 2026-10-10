"""Deterministic first-sheet extraction for Excel invoice workbooks."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl
import xlrd

from app.extraction.base import BaseExtractor
from app.extraction.exceptions import EmptyDocumentError, MalformedDocumentError
from app.extraction.models import DocumentFormat


class ExcelExtractor(BaseExtractor):
    @property
    def supported_format(self) -> DocumentFormat:
        return DocumentFormat.EXCEL

    @property
    def method_identifier(self) -> str:
        return "deterministic_excel"

    def _extract_impl(self, content: bytes | str) -> tuple[Any | None, Any | None, list[str], dict[str, Any]]:
        if not isinstance(content, bytes) or not content:
            raise EmptyDocumentError("Excel workbook is empty")

        try:
            if content.startswith(b"PK\x03\x04"):
                rows = self._read_xlsx(content)
            elif content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                rows = self._read_xls(content)
            else:
                raise MalformedDocumentError("File is not a supported .xls or .xlsx workbook")
        except MalformedDocumentError:
            raise
        except Exception as exc:
            raise MalformedDocumentError("Could not read the first worksheet in the Excel workbook") from exc

        if not rows:
            raise MalformedDocumentError("Excel worksheet must contain a header row and invoice line rows")

        return rows, None, [], {
            "table_count": 1,
            "character_count": sum(len(str(value)) for row in rows for value in row.values()),
            "field_confidence": {"header_integrity": 1.0, "format_confidence": 0.95},
        }

    def _read_xlsx(self, content: bytes) -> list[dict[str, Any]]:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            if not workbook.worksheets:
                raise MalformedDocumentError("Excel workbook contains no worksheets")
            sheet = workbook.worksheets[0]
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            return self._rows_from_values(headers, values)
        finally:
            workbook.close()

    def _read_xls(self, content: bytes) -> list[dict[str, Any]]:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        if workbook.nsheets < 1:
            raise MalformedDocumentError("Excel workbook contains no worksheets")
        sheet = workbook.sheet_by_index(0)
        values = (
            tuple(self._xls_value(sheet.cell(row_index, column_index), workbook.datemode)
                  for column_index in range(sheet.ncols))
            for row_index in range(sheet.nrows)
        )
        headers = next(values, None)
        rows = self._rows_from_values(headers, values)
        workbook.release_resources()
        return rows

    def _xls_value(self, cell: xlrd.sheet.Cell, datemode: int) -> Any:
        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return None
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        return cell.value

    def _rows_from_values(self, header_values: tuple[Any, ...] | None, values: Any) -> list[dict[str, Any]]:
        if not header_values:
            return []

        headers: list[str] = []
        seen: dict[str, int] = {}
        for index, value in enumerate(header_values, start=1):
            header = str(value).strip().lstrip("\ufeff") if value is not None else ""
            header = header or f"column_{index}"
            seen[header] = seen.get(header, 0) + 1
            headers.append(header if seen[header] == 1 else f"{header}_{seen[header]}")

        rows = []
        for values_row in values:
            row = {
                header: self._cell_value(values_row[index]) if index < len(values_row) else None
                for index, header in enumerate(headers)
            }
            if any(value not in (None, "") for value in row.values()):
                rows.append(row)
        return rows

    def _cell_value(self, value: Any) -> Any:
        if isinstance(value, datetime):
            return value.isoformat() if value.time().isoformat() != "00:00:00" else value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return value