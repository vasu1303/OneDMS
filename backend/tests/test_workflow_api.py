import asyncio
from datetime import date, datetime, timezone
from decimal import Decimal
import unittest
from unittest.mock import AsyncMock, MagicMock, patch

import httpx
from httpx import ASGITransport

from app.api.deps.db import get_db, get_db_session
from app.api.deps.storage import get_storage
from app.main import app
from app.models.invoice import (
    Dealer, DmsSystem, InboundDocument, InvoiceLineItem, MappingConfig,
    StandardizedInvoice, ValidationRule,
)
from app.services.oem import OEMExportService


class TestWorkflowAPI(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.dealer = Dealer(
            id=1,
            dealer_code="DEALER-TEST",
            name="Test Dealer Automotive",
            gstin="27AABCU9603R1ZM",
        )
        self.dms = DmsSystem(
            id=1,
            name="Test DMS Tier 1",
            integration_tier=1,
            integration_method="API",
            input_format="JSON",
        )
        self.document = InboundDocument(
            id=1,
            dealer_id=1,
            dms_id=1,
            document_type="INVOICE",
            received_at=datetime.now(timezone.utc),
            status="RECEIVED",
            raw_payload={
                "invoiceNo": "INV-TEST-001",
                "billDate": "2026-10-01",
                "dealerCode": "DEALER-TEST",
                "buyerCode": "DAIMLER-DEMO",
                "currencyCode": "INR",
                "subTotal": "1000.00",
                "taxTotal": "180.00",
                "grandTotal": "1180.00",
                "items": [
                    {
                        "sku": "PART-001",
                        "description": "Brake rotor",
                        "qty": "2.000",
                        "rate": "500.00",
                        "taxableValue": "1000.00",
                        "taxAmount": "180.00",
                        "lineTotal": "1180.00",
                        "category": "PART",
                    }
                ],
            },
        )
        self.document.dealer = self.dealer
        self.document.dms = self.dms
        self.document.invoice = None

        self.invoice = StandardizedInvoice(
            id=1,
            document_id=1,
            invoice_number="INV-TEST-001",
            invoice_date=date(2026, 10, 1),
            buyer_oem_id="DAIMLER-DEMO",
            currency="INR",
            subtotal=Decimal("1000.00"),
            tax_amount=Decimal("180.00"),
            total_amount=Decimal("1180.00"),
            canonical_payload={
                "invoice_number": "INV-TEST-001",
                "invoice_date": "2026-10-01",
                "supplier_dealer_code": "DEALER-TEST",
                "buyer_oem_id": "DAIMLER-DEMO",
                "currency": "INR",
                "subtotal": "1000.00",
                "tax_amount": "180.00",
                "total_amount": "1180.00",
                "line_items": [
                    {
                        "line_number": 1,
                        "item_code": "PART-001",
                        "description": "Brake rotor",
                        "quantity": "2.000",
                        "unit_price": "500.00",
                        "discount_amount": "0.00",
                        "taxable_amount": "1000.00",
                        "tax_rate": "18.00",
                        "tax_amount": "180.00",
                        "line_total": "1180.00",
                        "chassis_number": None,
                        "item_category": "PART",
                    }
                ],
                "_onedms": {
                    "extraction": {"confidence_score": 0.98, "method": "json_structured"},
                    "validation": {"status": "VALID", "issues": []},
                },
            },
            validation_status="VALID",
            review_status="NOT_REQUIRED",
            oem_delivery_status="NOT_SENT",
            created_at=datetime.now(timezone.utc),
            updated_at=datetime.now(timezone.utc),
        )
        line = InvoiceLineItem(
            id=1,
            invoice_id=1,
            line_number=1,
            item_code="PART-001",
            description="Brake rotor",
            quantity=Decimal("2.000"),
            unit_price=Decimal("500.00"),
            discount_amount=Decimal("0.00"),
            taxable_amount=Decimal("1000.00"),
            tax_rate=Decimal("18.00"),
            tax_amount=Decimal("180.00"),
            line_total=Decimal("1180.00"),
            chassis_number=None,
        )
        self.invoice.line_items = [line]
        self.invoice.document = self.document

    def _setup_dependencies(self, mock_session):
        mock_storage = MagicMock()
        mock_storage.check = AsyncMock(return_value=None)
        mock_storage.get_file_bytes = AsyncMock(return_value=(b"{}", "application/json"))

        async def override_db():
            yield mock_session

        app.dependency_overrides[get_db] = override_db
        app.dependency_overrides[get_db_session] = override_db
        app.dependency_overrides[get_storage] = lambda: mock_storage
        return mock_storage

    async def test_health_liveness(self):
        transport = ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            res = await client.get("/api/health/")
            self.assertEqual(res.status_code, 200)
            data = res.json()
            self.assertEqual(data["status"], "healthy")

    async def test_health_dependencies(self):
        transport = ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            with patch("app.api.v1.endpoints.health.check_database", AsyncMock(return_value=None)):
                app.state.database = MagicMock()
                mock_storage = MagicMock()
                mock_storage.check = AsyncMock(return_value=None)
                app.state.object_storage = mock_storage
                app.state.infrastructure_errors = {}

                res = await client.get("/api/health/dependencies")
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertEqual(data["status"], "healthy")
                self.assertEqual(data["database"]["status"], "healthy")
                self.assertEqual(data["object_storage"]["status"], "healthy")
                self.assertIn("llm", data)

    async def test_list_documents(self):
        mock_session = AsyncMock()
        mock_docs_res = MagicMock()
        mock_docs_res.scalars.return_value.all.return_value = [self.document]
        mock_session.execute.return_value = mock_docs_res

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.get("/api/documents")
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertIn("items", data)
                self.assertEqual(len(data["items"]), 1)
                item = data["items"][0]
                self.assertEqual(item["id"], 1)
                self.assertEqual(item["dealer_code"], "DEALER-TEST")
                self.assertEqual(item["processing_status"], "received")
        finally:
            app.dependency_overrides.clear()

    async def test_get_document_detail(self):
        mock_session = AsyncMock()
        mock_res = MagicMock()
        mock_res.scalar_one_or_none.return_value = self.document
        mock_session.execute.return_value = mock_res

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.get("/api/documents/1")
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertEqual(data["id"], 1)
                self.assertEqual(data["processing_status"], "received")
                self.assertEqual(data["source_type"], "json")
        finally:
            app.dependency_overrides.clear()

    async def test_process_document_success(self):
        fresh_doc = InboundDocument(
            id=1,
            dealer_id=1,
            dms_id=1,
            document_type="INVOICE",
            received_at=datetime.now(timezone.utc),
            status="RECEIVED",
            raw_payload=self.document.raw_payload,
        )
        fresh_doc.dealer = self.dealer
        fresh_doc.dms = self.dms
        fresh_doc.invoice = None

        mock_session = AsyncMock()
        mock_doc_res = MagicMock()
        mock_doc_res.scalar_one_or_none.return_value = fresh_doc

        mock_map_res = MagicMock()
        mock_map_res.scalar_one_or_none.return_value = None

        mock_rules_res = MagicMock()
        mock_rules_res.scalars.return_value.all.return_value = []

        mock_session.execute.side_effect = [mock_doc_res, mock_map_res, mock_rules_res]
        mock_session.add = MagicMock()
        mock_session.commit = AsyncMock()
        mock_session.flush = AsyncMock()
        mock_session.refresh = AsyncMock()

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.post("/api/documents/1/process")
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertEqual(data["id"], 1)
                self.assertEqual(data["processing_status"], "processed")
                self.assertEqual(data["validation_status"], "valid")
                self.assertEqual(data["review_status"], "pending")
        finally:
            app.dependency_overrides.clear()

    async def test_process_csv_and_excel_rows_use_profile_mapping(self):
        from io import BytesIO
        from openpyxl import Workbook

        headers = [
            "InvoiceNumber", "InvoiceDate", "DealerCode", "BuyerCode", "Currency",
            "Subtotal", "TaxTotal", "GrandTotal", "SKU", "Description", "Qty",
            "UnitPrice", "TaxableValue", "TaxAmount", "LineTotal",
        ]
        values = [
            "ROW-001", "2026-10-01", "DEALER-TEST", "DAIMLER-DEMO", "INR",
            "1000.00", "180.00", "1180.00", "PART-001", "Brake rotor", "2",
            "500.00", "1000.00", "180.00", "1180.00",
        ]
        csv_content = (",".join(headers) + "\n" + ",".join(values)).encode()
        workbook_buffer = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        sheet.append([values[0], values[1], values[2], values[3], values[4], 1000, 180, 1180,
                      values[8], values[9], 2, 500, 1000, 180, 1180])
        workbook.save(workbook_buffer)
        excel_content = workbook_buffer.getvalue()
        mapping_config = {
            "invoice_number": "InvoiceNumber", "invoice_date": "InvoiceDate",
            "supplier_dealer_code": "DealerCode", "buyer_oem_id": "BuyerCode",
            "currency": "Currency", "subtotal": "Subtotal", "tax_amount": "TaxTotal",
            "total_amount": "GrandTotal",
            "line_items": {
                "source_field": "rows", "item_code": "SKU", "description": "Description",
                "quantity": "Qty", "unit_price": "UnitPrice", "taxable_amount": "TaxableValue",
                "tax_amount": "TaxAmount", "line_total": "LineTotal",
            },
        }

        for input_format, filename, mime_type, content in (
            ("CSV", "invoice.csv", "text/csv", csv_content),
            ("EXCEL", "invoice.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", excel_content),
        ):
            with self.subTest(input_format=input_format):
                fresh_doc = InboundDocument(
                    id=1, dealer_id=1, dms_id=1, document_type="INVOICE",
                    received_at=datetime.now(timezone.utc), status="RECEIVED",
                    original_file_name=filename, mime_type=mime_type, storage_key="inbound/source",
                )
                fresh_doc.dealer = self.dealer
                fresh_doc.dms = DmsSystem(
                    id=1, name="Tabular DMS", integration_tier=2,
                    integration_method="UPLOAD", input_format=input_format,
                )
                fresh_doc.invoice = None
                mock_session = AsyncMock()
                doc_result = MagicMock()
                doc_result.scalar_one_or_none.return_value = fresh_doc
                mapping_result = MagicMock()
                mapping_result.scalar_one_or_none.return_value = MappingConfig(mapping_config=mapping_config)
                rules_result = MagicMock()
                rules_result.scalars.return_value.all.return_value = []
                mock_session.execute.side_effect = [doc_result, mapping_result, rules_result]
                mock_session.add = MagicMock()
                mock_session.commit = AsyncMock()
                mock_session.flush = AsyncMock()
                mock_session.refresh = AsyncMock()
                mock_storage = self._setup_dependencies(mock_session)
                mock_storage.get_file_bytes.return_value = (content, mime_type)
                try:
                    async with httpx.AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
                        response = await client.post("/api/documents/1/process")
                    self.assertEqual(response.status_code, 200, response.text)
                    invoice = next(
                        call.args[0] for call in mock_session.add.call_args_list
                        if isinstance(call.args[0], StandardizedInvoice)
                    )
                    self.assertEqual(invoice.canonical_payload["invoice_number"], "ROW-001")
                    self.assertEqual(invoice.canonical_payload["line_items"][0]["item_code"], "PART-001")
                    self.assertEqual(response.json()["processing_status"], "processed")
                    self.assertEqual(response.json()["source_type"], "excel" if input_format == "EXCEL" else "csv")
                finally:
                    app.dependency_overrides.clear()

    async def test_process_document_idempotent_existing_invoice(self):
        doc_with_inv = InboundDocument(
            id=1,
            dealer_id=1,
            dms_id=1,
            document_type="INVOICE",
            received_at=datetime.now(timezone.utc),
            status="COMPLETED",
            raw_payload=self.document.raw_payload,
        )
        doc_with_inv.dealer = self.dealer
        doc_with_inv.dms = self.dms
        doc_with_inv.invoice = self.invoice

        mock_session = AsyncMock()
        mock_doc_res = MagicMock()
        mock_doc_res.scalar_one_or_none.return_value = doc_with_inv
        mock_session.execute.return_value = mock_doc_res

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.post("/api/documents/1/process", json={"force": False})
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertEqual(data["invoice_id"], 1)
                self.assertEqual(data["processing_status"], "processed")
        finally:
            app.dependency_overrides.clear()

    async def test_list_invoices(self):
        mock_session = AsyncMock()
        mock_invs_res = MagicMock()
        mock_invs_res.scalars.return_value.all.return_value = [self.invoice]
        mock_session.execute.return_value = mock_invs_res

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.get("/api/invoices")
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertIn("items", data)
                self.assertEqual(len(data["items"]), 1)
                inv = data["items"][0]
                self.assertEqual(inv["header"]["invoice_number"], "INV-TEST-001")
                self.assertEqual(inv["header"]["total_amount"], 1180.00)
        finally:
            app.dependency_overrides.clear()

    async def test_get_invoice_detail(self):
        mock_session = AsyncMock()
        mock_res = MagicMock()
        mock_res.scalar_one_or_none.return_value = self.invoice
        mock_session.execute.return_value = mock_res

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.get("/api/invoices/1")
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertEqual(data["id"], 1)
                self.assertEqual(data["header"]["invoice_number"], "INV-TEST-001")
                self.assertEqual(len(data["line_items"]), 1)
                self.assertEqual(data["line_items"][0]["part_number"], "PART-001")
        finally:
            app.dependency_overrides.clear()

    async def test_patch_invoice_corrections(self):
        mock_session = AsyncMock()
        mock_inv_res = MagicMock()
        mock_inv_res.scalar_one_or_none.return_value = self.invoice
        mock_rules_res = MagicMock()
        mock_rules_res.scalars.return_value.all.return_value = []
        mock_session.execute.side_effect = [mock_inv_res, mock_rules_res]
        mock_session.commit = AsyncMock()
        mock_session.refresh = AsyncMock()

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                patch_payload = {
                    "header": {
                        "invoice_number": "INV-TEST-CORRECTED",
                        "currency_code": "INR",
                        "subtotal_amount": 1000.00,
                        "tax_amount": 180.00,
                        "total_amount": 1180.00,
                    },
                    "line_items": [
                        {
                            "line_number": 1,
                            "part_number": "PART-001-FIXED",
                            "description": "Corrected description",
                            "quantity": 2,
                            "unit_price": 500.0,
                            "discount_amount": 0.0,
                            "tax_amount": 180.0,
                            "line_total": 1180.0,
                        }
                    ],
                }
                res = await client.patch("/api/invoices/1", json=patch_payload)
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertEqual(data["header"]["invoice_number"], "INV-TEST-CORRECTED")
                self.assertEqual(data["line_items"][0]["part_number"], "PART-001-FIXED")
        finally:
            app.dependency_overrides.clear()

    async def test_review_invoice_block_invalid_without_override(self):
        self.invoice.validation_status = "INVALID"
        mock_session = AsyncMock()
        mock_inv_res = MagicMock()
        mock_inv_res.scalar_one_or_none.return_value = self.invoice
        mock_session.execute.return_value = mock_inv_res

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.post("/api/invoices/1/review", json={"decision": "approve"})
                self.assertEqual(res.status_code, 422)
                self.assertIn("Cannot approve an INVALID invoice", res.json()["detail"])
        finally:
            app.dependency_overrides.clear()

    async def test_review_invoice_approve_success(self):
        self.invoice.validation_status = "VALID"
        mock_session = AsyncMock()
        mock_inv_res = MagicMock()
        mock_inv_res.scalar_one_or_none.return_value = self.invoice
        mock_session.execute.return_value = mock_inv_res
        mock_session.commit = AsyncMock()
        mock_session.refresh = AsyncMock()

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.post("/api/invoices/1/review", json={"decision": "approve", "notes": "Verified"})
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertEqual(data["review_status"], "approved")
        finally:
            app.dependency_overrides.clear()

    async def test_mock_oem_export_preview(self):
        mock_session = AsyncMock()
        mock_inv_res = MagicMock()
        mock_inv_res.scalar_one_or_none.return_value = self.invoice
        mock_map_res = MagicMock()
        mock_map_res.scalar_one_or_none.return_value = None
        mock_session.execute.side_effect = [mock_inv_res, mock_map_res]

        self._setup_dependencies(mock_session)
        try:
            transport = ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                res = await client.post("/api/invoices/1/oem-preview")
                self.assertEqual(res.status_code, 200)
                data = res.json()
                self.assertTrue(data["is_mock"])
                self.assertEqual(data["target_system"], "DAIMLER_ERP")
                payload = data["payload"]
                self.assertNotIn("_onedms", payload)
                self.assertEqual(payload["invoiceNumber"], "INV-TEST-001")
        finally:
            app.dependency_overrides.clear()


if __name__ == "__main__":
    unittest.main()
